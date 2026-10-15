"""Regenerate the synthetic two-sheet fixture; never reads business data."""
import csv
from datetime import datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

fixture_dir = Path(__file__).resolve().parent
book = Workbook()
sales = book.active
sales.title = "ventas"
with (fixture_dir / "sample_sales.csv").open(encoding="utf-8", newline="") as source:
    rows = list(csv.reader(source))
sales.append(rows[0])
for row in rows[1:]:
    sales.append([int(row[0]), int(row[1]), int(row[2]), row[3], row[4] == "true", row[5] or None])
book.create_sheet("segunda").append(["marker"])
book["segunda"].append(["second-sheet"])
book.properties.created = book.properties.modified = datetime(2026, 10, 15)
buffer = BytesIO()
book.save(buffer)
target = fixture_dir / "two_sheets.xlsx"
with ZipFile(buffer) as original, ZipFile(target, "w") as output:
    for info in original.infolist():
        info.date_time = (2026, 10, 15, 0, 0, 0)
        data = original.read(info.filename)
        # openpyxl sets modified at save time; keep the fixture reproducible.
        if info.filename == "docProps/core.xml":
            import re
            data = re.sub(rb"(<dcterms:modified[^>]*>).*?(</dcterms:modified>)", rb"\g<1>2026-10-15T00:00:00Z\g<2>", data)
        output.writestr(info, data)
check = load_workbook(target, read_only=True)
assert check.sheetnames == ["ventas", "segunda"]
assert check["ventas"].max_row == 13
assert check["ventas"]["C2"].value == 0
assert check["segunda"]["A2"].value == "second-sheet"
check.close()
print("two_sheets.xlsx: ventas=12 filas, segunda=second-sheet")
